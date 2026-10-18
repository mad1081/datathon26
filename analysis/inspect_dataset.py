"""Audit local labels and MOV container metadata without transmitting video pixels."""
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import re
import struct

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def atoms(stream, start, end):
    pos = start
    while pos + 8 <= end:
        stream.seek(pos)
        size, kind = struct.unpack(">I4s", stream.read(8))
        header = 8
        if size == 1:
            size = struct.unpack(">Q", stream.read(8))[0]
            header = 16
        elif size == 0:
            size = end - pos
        if size < header or pos + size > end:
            raise ValueError(f"Invalid MOV atom at {pos}")
        yield kind, pos + header, pos + size
        pos += size


def payload(stream, atom):
    _, start, end = atom
    stream.seek(start)
    return stream.read(end - start)


def media_time(data):
    if data[0] == 1:
        scale, duration = struct.unpack_from(">IQ", data, 20)
    else:
        scale, duration = struct.unpack_from(">II", data, 12)
    return scale, duration


def inspect_mov(path):
    result = {"bytes": path.stat().st_size, "tracks": []}
    with path.open("rb") as stream:
        top = list(atoms(stream, 0, result["bytes"]))
        moov = next(a for a in top if a[0] == b"moov")
        children = list(atoms(stream, moov[1], moov[2]))
        mvhd = payload(stream, next(a for a in children if a[0] == b"mvhd"))
        scale, duration = media_time(mvhd)
        result["duration_sec"] = duration / scale
        for trak in (a for a in children if a[0] == b"trak"):
            track = {}
            tc = list(atoms(stream, trak[1], trak[2]))
            tkhd = payload(stream, next(a for a in tc if a[0] == b"tkhd"))
            width, height = struct.unpack_from(">II", tkhd, len(tkhd) - 8)
            track["display_width"] = width / 65536
            track["display_height"] = height / 65536
            matrix_offset = 52 if tkhd[0] == 1 else 40
            track["transform_matrix_raw"] = list(struct.unpack_from(">9i", tkhd, matrix_offset))
            mdia = next(a for a in tc if a[0] == b"mdia")
            mc = list(atoms(stream, mdia[1], mdia[2]))
            hdlr = payload(stream, next(a for a in mc if a[0] == b"hdlr"))
            track["handler"] = hdlr[8:12].decode("ascii", errors="replace")
            scale, duration = media_time(payload(stream, next(a for a in mc if a[0] == b"mdhd")))
            track["duration_sec"] = duration / scale
            minf = next(a for a in mc if a[0] == b"minf")
            stbl = next(a for a in atoms(stream, minf[1], minf[2]) if a[0] == b"stbl")
            sc = list(atoms(stream, stbl[1], stbl[2]))
            stsd = payload(stream, next(a for a in sc if a[0] == b"stsd"))
            track["codec"] = stsd[12:16].decode("ascii", errors="replace")
            if track["handler"] == "vide":
                track["coded_width"], track["coded_height"] = struct.unpack_from(">HH", stsd, 40)
                stts = payload(stream, next(a for a in sc if a[0] == b"stts"))
                entries = struct.unpack_from(">I", stts, 4)[0]
                timings = [struct.unpack_from(">II", stts, 8 + 8 * i) for i in range(entries)]
                track["frame_count"] = sum(n for n, delta in timings)
                ticks = sum(n * delta for n, delta in timings)
                track["average_fps"] = track["frame_count"] * scale / ticks
                track["sample_duration_variants"] = len({delta for n, delta in timings})
            result["tracks"].append(track)
    with path.open("rb") as stream:
        result["sha256"] = hashlib.file_digest(stream, "sha256").hexdigest()
    return result


def parse_timestamp(raw):
    if raw is None:
        return None
    text = str(raw)
    minutes = re.search(r"(\d+)\s*мин", text)
    seconds = re.search(r"(\d+)\s*сек", text)
    if not minutes and not seconds:
        raise ValueError(f"Unrecognized timestamp: {raw}")
    return (int(minutes[1]) * 60 if minutes else 0) + (int(seconds[1]) if seconds else 0)


workbook = load_workbook(ROOT / "open/open_labels.xlsx", read_only=True, data_only=True)
records = []
issues = []
for row in workbook["labels"].iter_rows(min_row=2, values_only=True):
    name, timestamp, label = row[:3]
    if name is None:
        continue
    path = ROOT / "open/videos" / f"{name}.mov"
    record = {"file": path.name, "label": label, "timestamp_raw": timestamp,
              "timestamp_sec": parse_timestamp(timestamp)}
    if path.exists():
        record.update(inspect_mov(path))
        if record["timestamp_sec"] is not None and not 0 <= record["timestamp_sec"] <= record["duration_sec"]:
            issues.append({"file": path.name, "issue": "timestamp outside duration"})
    else:
        issues.append({"file": path.name, "issue": "missing video"})
    records.append(record)
workbook.close()
groups = defaultdict(list)
for record in records:
    groups[record.get("sha256", record["file"])].append(record)
duplicates = [{"files": [r["file"] for r in group], "labels": [r["label"] for r in group],
               "timestamps_sec": [r["timestamp_sec"] for r in group]}
              for group in groups.values() if len(group) > 1]
videos = [track for r in records for track in r.get("tracks", []) if track["handler"] == "vide"]
actual_files = {p.name for p in (ROOT / "open/videos").glob("*.mov")}
summary = {
    "video_files": len(actual_files), "label_rows": len(records),
    "labels": dict(Counter(r["label"] for r in records)),
    "unique_files_by_sha256": len(groups), "duplicate_groups": duplicates,
    "unlabeled_files": sorted(actual_files - {r["file"] for r in records}),
    "issues": issues,
    "duration_range_sec": [min(r["duration_sec"] for r in records), max(r["duration_sec"] for r in records)],
    "total_duration_sec": sum(r["duration_sec"] for r in records),
    "file_size_range_bytes": [min(r["bytes"] for r in records), max(r["bytes"] for r in records)],
    "total_video_bytes": sum(r["bytes"] for r in records),
    "video_codecs": dict(Counter(t["codec"] for t in videos)),
    "coded_resolutions": dict(Counter(f'{t["coded_width"]}x{t["coded_height"]}' for t in videos)),
    "average_fps_range": [min(t["average_fps"] for t in videos), max(t["average_fps"] for t in videos)],
    "videos_with_audio_tracks": [r["file"] for r in records if any(t["handler"] == "soun" for t in r["tracks"])],
    "videos_with_multiple_sample_durations": sum(t["sample_duration_variants"] > 1 for t in videos),
    "method": "Read all label rows, SHA-256 hashes and MOV metadata; no frames decoded or transmitted.",
}
output = ROOT / "analysis/dataset_audit.json"
output.write_text(json.dumps({"summary": summary, "records": records}, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
